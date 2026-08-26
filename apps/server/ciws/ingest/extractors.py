"""Getting text out of whatever you drop on the hub.

Every heavy parser is an optional import resolved at call time. That is the
whole design: a fresh install with nothing but the base dependencies still
ingests Markdown, text, code, CSV, JSON and HTML, and tells you exactly which
package to install for the format it cannot read yet -- instead of failing at
import time and taking the server down with it.
"""

from __future__ import annotations

import csv
import io
import json
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from ..core.errors import ValidationFailed
from ..core.logging import get_logger

log = get_logger("ingest.extract")

#: A single document is capped here. Past a few megabytes of text the useful
#: move is to split the source, not to chew through it in one transaction.
MAX_TEXT_CHARS = 12_000_000

TEXT_SUFFIXES = {
    ".txt", ".md", ".markdown", ".rst", ".log", ".text", ".adoc", ".org",
}
CODE_SUFFIXES = {
    ".py": "python", ".js": "javascript", ".jsx": "javascript", ".ts": "typescript",
    ".tsx": "typescript", ".java": "java", ".go": "go", ".rs": "rust", ".c": "c",
    ".cpp": "cpp", ".cc": "cpp", ".h": "c", ".hpp": "cpp", ".sh": "bash",
    ".bash": "bash", ".zsh": "bash", ".sql": "sql", ".rb": "ruby", ".php": "php",
    ".cs": "csharp", ".kt": "kotlin", ".swift": "swift", ".scala": "scala",
    ".r": "r", ".lua": "lua", ".pl": "perl", ".vue": "vue", ".svelte": "svelte",
    ".toml": "toml", ".ini": "ini", ".cfg": "ini", ".xml": "xml", ".gradle": "gradle",
    ".dockerfile": "docker", ".tf": "terraform", ".proto": "protobuf",
}
DATA_SUFFIXES = {".json", ".jsonl", ".ndjson", ".yaml", ".yml"}
TABLE_SUFFIXES = {".csv", ".tsv"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tiff"}

MIME_BY_SUFFIX = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".html": "text/html", ".htm": "text/html", ".md": "text/markdown",
    ".csv": "text/csv", ".json": "application/json",
}


@dataclass(slots=True)
class Extracted:
    text: str = ""
    title: str = ""
    page_count: int = 0
    mime_type: str = ""
    #: Character offsets where each page begins, so chunks can cite a page.
    page_offsets: list[int] = field(default_factory=list)
    meta: dict[str, Any] = field(default_factory=dict)


def supported_extensions() -> set[str]:
    return (
        TEXT_SUFFIXES
        | set(CODE_SUFFIXES)
        | DATA_SUFFIXES
        | TABLE_SUFFIXES
        | IMAGE_SUFFIXES
        | {".pdf", ".docx", ".xlsx", ".xls", ".pptx", ".html", ".htm"}
    )


def _need(package: str, fmt: str) -> ValidationFailed:
    return ValidationFailed(
        f"Reading {fmt} needs the '{package}' package. "
        f"Install it with: pip install {package}",
        package=package,
        format=fmt,
    )


def _read_text(path: Path) -> str:
    raw = path.read_bytes()[: MAX_TEXT_CHARS * 4]
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def strip_html(html: str) -> str:
    """Readable text from HTML, with bs4 when available and regex when not."""
    try:
        from bs4 import BeautifulSoup

        soup = BeautifulSoup(html, "html.parser")
        for tag in soup(["script", "style", "nav", "footer", "noscript", "svg", "form"]):
            tag.decompose()
        text = soup.get_text("\n")
    except ImportError:
        text = re.sub(r"(?is)<(script|style|nav|footer|noscript|svg)\b.*?</\1>", " ", html)
        text = re.sub(r"(?i)<br\s*/?>", "\n", text)
        text = re.sub(r"(?i)</(p|div|h[1-6]|li|tr)>", "\n", text)
        text = re.sub(r"<[^>]+>", " ", text)
        text = (
            text.replace("&nbsp;", " ").replace("&amp;", "&").replace("&lt;", "<")
            .replace("&gt;", ">").replace("&quot;", '"').replace("&#39;", "'")
        )
    lines = [ln.strip() for ln in text.splitlines()]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(ln for ln in lines if ln))


def html_title(html: str) -> str:
    match = re.search(r"(?is)<title[^>]*>(.*?)</title>", html)
    return re.sub(r"\s+", " ", match.group(1)).strip()[:300] if match else ""


def _extract_pdf(path: Path) -> Extracted:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise _need("pypdf", "PDF") from exc

    reader = PdfReader(str(path))
    pages: list[str] = []
    offsets: list[int] = []
    cursor = 0
    for page in reader.pages:
        try:
            content = page.extract_text() or ""
        except Exception:  # noqa: BLE001 - a single malformed page should not lose the file
            content = ""
        offsets.append(cursor)
        pages.append(content)
        cursor += len(content) + 2

    text = "\n\n".join(pages)
    info = getattr(reader, "metadata", None) or {}
    title = str(info.get("/Title", "") or "").strip()
    scanned = len(text.strip()) < 40 * len(pages)
    return Extracted(
        text=text,
        title=title,
        page_count=len(pages),
        mime_type="application/pdf",
        page_offsets=offsets,
        meta={
            "author": str(info.get("/Author", "") or ""),
            "likely_scanned": scanned,
            **({"note": "Little extractable text -- this PDF is probably scanned images."} if scanned else {}),
        },
    )


def _extract_docx(path: Path) -> Extracted:
    try:
        import docx
    except ImportError as exc:
        raise _need("python-docx", "Word documents") from exc

    document = docx.Document(str(path))
    blocks = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                blocks.append(" | ".join(cells))
    core = document.core_properties
    return Extracted(
        text="\n\n".join(blocks),
        title=(core.title or "").strip(),
        mime_type=MIME_BY_SUFFIX[".docx"],
        meta={"author": core.author or "", "tables": len(document.tables)},
    )


def _extract_xlsx(path: Path) -> Extracted:
    try:
        import openpyxl
    except ImportError as exc:
        raise _need("openpyxl", "Excel workbooks") from exc

    book = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    parts: list[str] = []
    for sheet in book.worksheets:
        parts.append(f"## Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if c is None else str(c) for c in row]
            if any(c.strip() for c in cells):
                parts.append(" | ".join(cells))
        parts.append("")
    book.close()
    return Extracted(
        text="\n".join(parts),
        title=path.stem,
        mime_type=MIME_BY_SUFFIX[".xlsx"],
        meta={"sheets": len(book.sheetnames)},
    )


def _extract_pptx(path: Path) -> Extracted:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise _need("python-pptx", "PowerPoint decks") from exc

    deck = Presentation(str(path))
    parts: list[str] = []
    offsets: list[int] = []
    cursor = 0
    for i, slide in enumerate(deck.slides, 1):
        offsets.append(cursor)
        chunk = [f"## Slide {i}"]
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                chunk.append(shape.text_frame.text.strip())
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            chunk.append(f"[Notes] {slide.notes_slide.notes_text_frame.text.strip()}")
        block = "\n".join(chunk)
        parts.append(block)
        cursor += len(block) + 2
    return Extracted(
        text="\n\n".join(parts),
        title=path.stem,
        page_count=len(deck.slides),
        mime_type=MIME_BY_SUFFIX[".pptx"],
        page_offsets=offsets,
    )


def _extract_table(path: Path) -> Extracted:
    raw = _read_text(path)
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    try:
        rows = list(csv.reader(io.StringIO(raw), delimiter=delimiter))
    except csv.Error as exc:
        raise ValidationFailed(f"Could not parse {path.name}: {exc}") from exc

    lines = [" | ".join(r) for r in rows if any(c.strip() for c in r)]
    header = rows[0] if rows else []
    return Extracted(
        text="\n".join(lines),
        title=path.stem,
        mime_type="text/csv",
        meta={"rows": max(0, len(rows) - 1), "columns": header},
    )


def _extract_data(path: Path) -> Extracted:
    raw = _read_text(path)
    suffix = path.suffix.lower()
    if suffix in {".json", ".jsonl", ".ndjson"}:
        try:
            if suffix == ".json":
                text = json.dumps(json.loads(raw), indent=2, ensure_ascii=False)
            else:
                text = "\n".join(
                    json.dumps(json.loads(ln), ensure_ascii=False)
                    for ln in raw.splitlines()
                    if ln.strip()
                )
        except json.JSONDecodeError:
            text = raw  # malformed JSON is still searchable text
    else:
        text = raw
    return Extracted(text=text, title=path.stem, mime_type="application/json", meta={"format": suffix})


async def extract(path: Path | str) -> Extracted:
    """Read a file into text. Raises ValidationFailed with a fixable message."""
    path = Path(path)
    if not path.exists():
        raise ValidationFailed(f"No such file: {path}")
    if path.is_dir():
        raise ValidationFailed(f"{path} is a directory -- use ingest_directory")

    suffix = path.suffix.lower()
    size = path.stat().st_size

    if suffix == ".pdf":
        result = _extract_pdf(path)
    elif suffix == ".docx":
        result = _extract_docx(path)
    elif suffix in {".xlsx", ".xls"}:
        result = _extract_xlsx(path)
    elif suffix == ".pptx":
        result = _extract_pptx(path)
    elif suffix in {".html", ".htm"}:
        raw = _read_text(path)
        result = Extracted(
            text=strip_html(raw),
            title=html_title(raw) or path.stem,
            mime_type="text/html",
        )
    elif suffix in TABLE_SUFFIXES:
        result = _extract_table(path)
    elif suffix in DATA_SUFFIXES:
        result = _extract_data(path)
    elif suffix in IMAGE_SUFFIXES:
        # No OCR here. The pipeline can caption it with a vision model instead,
        # which handles diagrams and screenshots far better than OCR does.
        result = Extracted(
            text="",
            title=path.stem,
            mime_type=f"image/{suffix.lstrip('.')}",
            meta={"needs_vision": True},
        )
    elif suffix in CODE_SUFFIXES:
        result = Extracted(
            text=_read_text(path),
            title=path.name,
            mime_type="text/plain",
            meta={"language": CODE_SUFFIXES[suffix], "is_code": True},
        )
    elif suffix in TEXT_SUFFIXES or suffix == "":
        result = Extracted(text=_read_text(path), title=path.stem, mime_type="text/plain")
    else:
        # Unknown extension: try it as text rather than refusing outright.
        text = _read_text(path)
        printable = sum(1 for c in text[:4000] if c.isprintable() or c in "\n\r\t")
        if not text or printable < len(text[:4000]) * 0.85:
            raise ValidationFailed(
                f"Don't know how to read '{suffix or path.name}'. "
                f"Supported: {', '.join(sorted(supported_extensions()))}"
            )
        result = Extracted(text=text, title=path.stem, mime_type="text/plain")

    if len(result.text) > MAX_TEXT_CHARS:
        result.text = result.text[:MAX_TEXT_CHARS]
        result.meta["truncated"] = True
    result.title = (result.title or path.stem).strip()
    result.mime_type = result.mime_type or MIME_BY_SUFFIX.get(suffix, "application/octet-stream")
    result.meta.setdefault("size_bytes", size)
    result.meta.setdefault("filename", path.name)
    return result


async def extract_url(url: str, *, timeout: float = 25.0) -> Extracted:
    """Fetch a URL and reduce it to readable text."""
    import httpx

    headers = {
        # Plenty of sites serve a stub to unknown agents; a browser UA gets the article.
        "user-agent": (
            "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 "
            "(KHTML, like Gecko) Chrome/122.0 Safari/537.36"
        ),
        "accept": "text/html,application/xhtml+xml,application/pdf;q=0.9,*/*;q=0.8",
    }
    async with httpx.AsyncClient(timeout=timeout, follow_redirects=True, headers=headers) as client:
        response = await client.get(url)
        response.raise_for_status()
        content_type = response.headers.get("content-type", "").split(";")[0].strip()
        body = response.content

    if "pdf" in content_type:
        import tempfile

        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as handle:
            handle.write(body)
            temp = Path(handle.name)
        try:
            result = _extract_pdf(temp)
        finally:
            temp.unlink(missing_ok=True)
        result.title = result.title or url
        result.meta["source_url"] = url
        return result

    text = body.decode(response.encoding or "utf-8", errors="replace")
    if "html" in content_type or text.lstrip()[:200].lower().startswith(("<!doctype", "<html")):
        return Extracted(
            text=strip_html(text),
            title=html_title(text) or url,
            mime_type="text/html",
            meta={"source_url": url, "status": 200},
        )
    return Extracted(
        text=text, title=url, mime_type=content_type or "text/plain", meta={"source_url": url}
    )
